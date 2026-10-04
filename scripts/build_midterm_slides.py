"""Dựng slide báo cáo giữa kỳ (`report/slide.pptx`) từ template IUH và các bảng do script sinh.

Mọi số trong bảng đọc từ `docs/thesis/tables`, `report/tables` và `report/src` (giống DOCX);
ví dụ thật đọc từ artifact đã khóa. Không chạy suy luận hay đánh giá nào.

    .venv/Scripts/python.exe -m scripts.build_midterm_slides [--output report/slide.pptx]
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from scripts.midterm_slides import deck as deck_module
from scripts.midterm_slides import part_a, part_b, part_c, part_d

SLIDES = [
    part_a.intro, part_a.research_questions, part_a.system, part_a.demo, part_a.datasets,
    part_a.leakage, part_a.split_rules, part_a.rules, part_a.scenario_map,
    part_b.kb1, part_b.kb2, part_b.kb3, part_b.kb4, part_b.kb5, part_b.kb6,
    part_c.classifier, part_c.rq1, part_c.diagnosis, part_c.selection, part_c.progress,
    part_c.rq2, part_c.rq2_limits, part_c.rq3, part_c.parser_parity, part_c.summary,
    part_d.schedule, part_d.next_plan, part_d.questions,
]


def build(output: Path) -> deck_module.Deck:
    deck = deck_module.open_template()
    deck_module.cover(deck)
    for make in SLIDES:
        make(deck)
    if deck.count != deck_module.MAIN_TOTAL:
        raise RuntimeError(f"Có {deck.count} slide chính, footer ghi {deck_module.MAIN_TOTAL}")
    for key, page in part_a.RESULT_PAGES.items():
        if deck.pages.get(key) != page:
            raise RuntimeError(f"Slide kết quả '{key}' ở trang {deck.pages.get(key)}, "
                               f"các slide kịch bản trỏ tới {page}")
    part_d.appendix(deck)
    deck_module.save(deck, output)
    return deck


def write_outline(pptx: Path, output: Path) -> None:
    """Kịch bản trình bày (tiêu đề + ghi chú người nói) sinh từ chính tệp slide."""
    from pptx import Presentation

    lines = ["# Kịch bản trình bày slide giữa kỳ", "",
             "> Sinh tự động bởi `scripts.build_midterm_slides` từ `report/slide.pptx`; sửa nội "
             "dung trong `scripts/midterm_slides/`, không sửa tay tệp này.", ""]
    for index, slide in enumerate(Presentation(str(pptx)).slides, start=1):
        heading = next((shape.text_frame.text for shape in slide.shapes
                        if shape.name == "Title" and shape.has_text_frame), "")
        heading = heading or ("Bìa" if index == 1 else "Hỏi đáp")
        lines += [f"## Slide {index} — {heading}", ""]
        if slide.has_notes_slide:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                lines += [note, ""]
    output.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--output", type=Path,
                        default=deck_module.TEMPLATE.parent / "slide.pptx")
    parser.add_argument("--outline", type=Path,
                        default=deck_module.TEMPLATE.parent / "slide" / "kich_ban_slide_giua_ky.md")
    args = parser.parse_args()
    deck = build(args.output)
    write_outline(args.output, args.outline)
    print(f"Đã ghi {args.output}: {deck.count} slide chính + {deck.appendix} phụ lục; "
          f"kịch bản: {args.outline}")


if __name__ == "__main__":
    main()
