"""Khung bộ slide: mở template IUH, giữ bìa, xóa slide mẫu, tạo slide có tiêu đề và số trang."""

from __future__ import annotations

import copy
import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.oxml.slide import CT_NotesMaster

from scripts.midterm_slides import kit

TEMPLATE = kit.ROOT / "report" / "example_template.pptx"
FOOTER = "Báo cáo giữa kỳ KLTN · Phạm Hoàng Phát – Thân Trọng Tính · KHDL18A"
MAIN_TOTAL = 29
SECTIONS = {
    1: "Giới thiệu đề tài",
    2: "Hệ thống và luồng xử lý",
    3: "Dữ liệu",
    4: "Phương pháp và kịch bản",
    5: "Kết quả",
    6: "Tiến độ và kế hoạch",
}


@dataclass
class Deck:
    pres: Presentation
    count: int = 0
    appendix: int = 0
    pages: dict[str, int] = field(default_factory=dict)

    def layout(self, name: str):
        return next(layout for layout in self.pres.slide_layouts if layout.name == name)

    def content(self, section: int, heading: str, *, key: str | None = None):
        """Slide nội dung nền sáng có ô số phần, tiêu đề và số trang."""
        slide = self.pres.slides.add_slide(self.layout("IUH – Trơn"))
        self.count += 1
        if key:
            self.pages[key] = self.count
        footer = f"Phần {section} · {SECTIONS[section]}    {self.count} / {MAIN_TOTAL}"
        kit.title(slide, f"0{section}", heading, footer)
        return slide

    def appendix_slide(self, heading: str):
        slide = self.pres.slides.add_slide(self.layout("IUH – Trơn"))
        self.appendix += 1
        kit.title(slide, "PL", heading,
                  f"Phụ lục {self.appendix} · không trình bày, dùng khi được hỏi")
        return slide

    def dark(self):
        self.count += 1
        return self.pres.slides.add_slide(self.layout("IUH – Nền tối"))


def _repair_notes(pres: Presentation) -> None:
    """Notes master của template không có placeholder → ghi chú người trình bày không có chỗ.

    Chép các placeholder của notes master mặc định vào, và bỏ notes slide cũ của bìa (rỗng)
    để python-pptx tạo lại từ master đã sửa.
    """
    sp_tree = pres.notes_master._element.cSld.spTree
    if not sp_tree.findall(qn("p:sp")):
        for shape in CT_NotesMaster.new_default().cSld.spTree.findall(qn("p:sp")):
            sp_tree.append(copy.deepcopy(shape))
    cover_part = pres.slides[0].part
    for r_id, rel in list(cover_part.rels.items()):
        if rel.reltype == RT.NOTES_SLIDE:
            cover_part.drop_rel(r_id)


def open_template() -> Deck:
    pres = Presentation(str(TEMPLATE))
    slide_ids = pres.slides._sldIdLst
    for slide_id in list(slide_ids)[1:]:
        pres.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)
    _repair_notes(pres)
    for layout in pres.slide_layouts:
        for shape in layout.shapes:
            if shape.has_text_frame and "Tương tác Người" in shape.text_frame.text:
                runs = shape.text_frame.paragraphs[0].runs
                runs[0].text = FOOTER
                for extra in runs[1:]:
                    extra.text = ""
    return Deck(pres)


def cover(deck: Deck) -> None:
    """Giữ ảnh nền và thẻ của bìa template, thay toàn bộ chữ."""
    slide = deck.pres.slides[0]
    deck.count += 1
    for shape in list(slide.shapes):
        if shape.name not in {"Image 0", "Title card"}:
            shape._element.getparent().remove(shape._element)
    kit.box(slide, 0.86, 2.0, 2.25, 0.42, fill=kit.NAVY, line=None, rounded=True)
    kit.text(slide, 0.86, 2.0, 2.25, 0.42, "BÁO CÁO GIỮA KỲ", size=14, bold=True,
             color=kit.WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    kit.text(slide, 3.25, 2.0, 4.2, 0.42, "Khóa luận tốt nghiệp · Khoa học dữ liệu", size=13,
             bold=True, color=kit.GRAY, anchor=MSO_ANCHOR.MIDDLE)
    kit.text(slide, 0.86, 2.6, 6.55, 1.6,
             "Xây dựng hệ thống phát hiện, sinh mô tả và truy vấn sự kiện âm thanh môi "
             "trường dựa trên RAG", size=24, bold=True, color=kit.NAVY)
    kit.text(slide, 0.86, 4.2, 6.55, 0.6,
             "Grounded Environmental Sound Event Detection, Captioning, and RAG-based "
             "Retrieval", size=13, italic=True, color=kit.GRAY)
    kit.box(slide, 0.86, 4.92, 6.55, 0.02, fill=kit.LINE, line=None)
    kit.text(slide, 0.86, 5.08, 6.55, 1.6, [
        [kit.Run("Phạm Hoàng Phát – 22703861", bold=True, color=kit.NAVY)],
        [kit.Run("Thân Trọng Tính – 22665711", bold=True, color=kit.NAVY)],
        [kit.Run("GVHD: ", color=kit.GRAY), kit.Run("TS. Huỳnh Công Bằng", bold=True,
                                                    color=kit.NAVY)],
        [kit.Run("Lớp KHDL18A · Tháng 10/2026", color=kit.GRAY)],
    ], size=15, space=3)
    kit.notes(slide, """
Mở đầu (30 giây). Chào thầy, nhóm em gồm Phát và Tính, báo cáo giữa kỳ đề tài
"Xây dựng hệ thống phát hiện, sinh mô tả và truy vấn sự kiện âm thanh môi trường dựa trên RAG".
Bài trình bày có 6 phần; nhóm tập trung nhiều nhất vào phương pháp, kịch bản thí nghiệm và kết
quả.""")


def save(deck: Deck, path: Path) -> None:
    buffer = io.BytesIO()
    deck.pres.save(buffer)
    _fix_jpeg_content_type(buffer.getvalue(), path)


def _fix_jpeg_content_type(data: bytes, path: Path) -> None:
    """Template khai báo ảnh bìa là "image/jpg" (không chuẩn); thay bằng Default jpg image/jpeg."""
    with zipfile.ZipFile(io.BytesIO(data)) as source, \
            zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            payload = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                xml = payload.decode("utf-8")
                xml = re.sub(r'<Override PartName="[^"]+\.jpg" ContentType="image/jpg"/>', "", xml)
                if 'Extension="jpg"' not in xml:
                    xml = xml.replace("<Default ", '<Default Extension="jpg" '
                                      'ContentType="image/jpeg"/><Default ', 1)
                payload = xml.encode("utf-8")
            target.writestr(item, payload)
