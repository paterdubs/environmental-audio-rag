"""Thành phần vẽ dùng chung: tiêu đề, thẻ, chữ có gạch đầu dòng, bảng, ảnh, biểu đồ."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from scripts.report_docx.markdown import USED_MARK, expand_include

ROOT = Path(__file__).resolve().parents[2]
FONT = "Arial"

NAVY = "062F5C"
BLUE = "1677B8"
YELLOW = "F4B336"
GREEN = "2E9B67"
RED = "B42318"
GRAY = "667085"
TEXT = "17212B"
LINE = "DFE5EC"
WHITE = "FFFFFF"
USED_FILL = "D9D9D9"
READ_FILL = "EAF2FB"
WARN_FILL = "FFF4DB"
WARN_TEXT = "8A5A00"
SOFT = "F0F4F9"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


@dataclass
class Run:
    text: str
    bold: bool = False
    color: str | None = None
    italic: bool = False


Para = str | list[Run]


def _runs(para: Para) -> list[Run]:
    if isinstance(para, str):
        para = para.replace("`", "")
        return [Run(piece[2:-2], bold=True) if piece.startswith("**") else Run(piece)
                for piece in re.split(r"(\*\*[^*]+\*\*)", para) if piece]
    return para


def _bullet(paragraph, color: str = BLUE) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    ppr.set("marL", str(Inches(0.24)))
    ppr.set("indent", str(-Inches(0.2)))
    clr = etree.SubElement(ppr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", color)
    etree.SubElement(ppr, qn("a:buFont")).set("typeface", FONT)
    etree.SubElement(ppr, qn("a:buChar")).set("char", "•")


def _items(paras) -> list[Para]:
    if isinstance(paras, str):
        return [paras]
    if paras and isinstance(paras[0], Run):
        return [paras]
    return list(paras)


def text(slide, x, y, w, h, paras, *, size=15, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, bullets=False, space=4, name=None, italic=False, margin=0.0):
    """Hộp chữ; mỗi đoạn là chuỗi (có thể chứa **đậm**) hoặc danh sách Run."""
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shape.name = name
    frame = shape.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(frame, side, Inches(margin))
    for index, para in enumerate(_items(paras)):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(space)
        if bullets:
            _bullet(paragraph)
        for piece in _runs(para):
            run = paragraph.add_run()
            run.text = piece.text
            font = run.font
            font.name = FONT
            font.size = Pt(size)
            font.bold = bold or piece.bold
            font.italic = italic or piece.italic
            font.color.rgb = rgb(piece.color or color)
    return shape


def box(slide, x, y, w, h, fill=WHITE, line=LINE, *, rounded=False, name=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shape.adjustments[0] = 0.06
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    if name:
        shape.name = name
    return shape


def badge(slide, x, y, label, *, size=0.56, fill=YELLOW, color=NAVY, font=18, oval=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL if oval else MSO_SHAPE.ROUNDED_RECTANGLE,
                                   Inches(x), Inches(y), Inches(size), Inches(size))
    if not oval:
        shape.adjustments[0] = 0.18
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background()
    shape.shadow.inherit = False
    frame = shape.text_frame
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(frame, side, 0)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = label
    run.font.name = FONT
    run.font.size = Pt(font)
    run.font.bold = True
    run.font.color.rgb = rgb(color)
    return shape


def title(slide, number: str, heading: str, footer: str) -> None:
    badge(slide, 0.5, 0.42, number)
    text(slide, 1.22, 0.3, 11.6, 0.82, heading, size=24, bold=True, color=NAVY,
         anchor=MSO_ANCHOR.MIDDLE, name="Title")
    text(slide, 7.3, 7.05, 5.53, 0.28, footer, size=10, color=GRAY, align=PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, heading, body, *, accent=BLUE, fill=WHITE, size=14, bullets=True,
         head_size=14, line=LINE):
    """Thẻ có nhãn đầu màu và nội dung."""
    box(slide, x, y, w, h, fill=fill, line=line, rounded=True)
    top = y + 0.12
    if heading:
        text(slide, x + 0.18, top, w - 0.36, 0.36, heading, size=head_size, bold=True,
             color=accent)
        top += 0.42
    text(slide, x + 0.18, top, w - 0.36, y + h - top - 0.08, body, size=size, bullets=bullets)


def callouts(slide, x, y, w, h, read, warn, *, size=13, split=0.6, horizontal=False):
    """Hai ô cố định cho slide kết quả: 'Đọc kết quả' và 'Không có nghĩa là'."""
    gap = 0.18
    if horizontal:
        read_w = (w - gap) * split
        boxes = [(x, y, read_w, h), (x + read_w + gap, y, w - read_w - gap, h)]
    else:
        read_h = (h - gap) * split
        boxes = [(x, y, w, read_h), (x, y + read_h + gap, w, h - read_h - gap)]
    card(slide, *boxes[0], "Đọc kết quả", read, accent=BLUE, fill=READ_FILL, line=READ_FILL,
         size=size)
    card(slide, *boxes[1], "Không có nghĩa là", warn, accent=WARN_TEXT, fill=WARN_FILL,
         line=WARN_FILL, size=size)


def frame_row(slide, cells: list[tuple[str, str]], y=1.3, h=1.12, size=12.5):
    """Khung kịch bản 4 ô: Câu hỏi · Cách làm · So sánh · Đo bằng."""
    gap = 0.15
    width = (12.33 - gap * (len(cells) - 1)) / len(cells)
    colors = [NAVY, BLUE, GREEN, WARN_TEXT]
    for index, (head, body) in enumerate(cells):
        x = 0.5 + index * (width + gap)
        box(slide, x, y, width, h, fill=WHITE, line=LINE, rounded=True)
        badge(slide, x + 0.14, y + 0.13, str(index + 1), size=0.32, fill=colors[index],
              color=WHITE, font=11, oval=True)
        text(slide, x + 0.54, y + 0.1, width - 0.64, 0.36, head, size=12.5, bold=True,
             color=colors[index], anchor=MSO_ANCHOR.MIDDLE)
        text(slide, x + 0.16, y + 0.5, width - 0.3, h - 0.56, body, size=size)


def picture(slide, path: Path, x, y, w, h, *, align="center"):
    """Ảnh vừa khung, giữ tỷ lệ."""
    with Image.open(path) as image:
        ratio = image.width / image.height
    if w / h > ratio:
        pic_h, pic_w = h, h * ratio
    else:
        pic_w, pic_h = w, w / ratio
    left = x + (w - pic_w) / 2
    top = y + ((h - pic_h) / 2 if align == "center" else 0)
    return slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(pic_w),
                                    Inches(pic_h))


def notes(slide, body: str) -> None:
    slide.notes_slide.notes_text_frame.text = body.strip()


def arrow(slide, x, y, w=0.36, h=0.3, color=GRAY):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w),
                                   Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def stat(slide, x, y, w, h, value, label, *, color=NAVY, fill=WHITE, value_size=28, size=12.5):
    box(slide, x, y, w, h, fill=fill, line=LINE, rounded=True)
    text(slide, x + 0.18, y + 0.08, w - 0.36, 0.6, value, size=value_size, bold=True,
         color=color)
    text(slide, x + 0.18, y + 0.7, w - 0.36, h - 0.76, label, size=size)


# ---------------------------------------------------------------- bảng
@dataclass
class TableData:
    header: list[str]
    rows: list[list[str]]
    used: list[bool]
    source: str


def load_table(rel_path: str, options: str = "") -> TableData:
    """Đọc bảng do script sinh, qua cùng bộ lọc trình bày với báo cáo DOCX."""
    lines = expand_include(ROOT / rel_path, options)
    rows = [[c.strip() for c in line.strip().strip("|").split("|")]
            for line in lines if line.startswith("|")]
    body = rows[2:]
    used = [row[0].startswith(USED_MARK) for row in body]
    body = [[c.replace(USED_MARK, "") for c in row] for row in body]
    note = next((line for line in lines if line.startswith("@@NOTE:")), "")
    source = re.sub(r"^@@NOTE:Nguồn:\s*", "", note).split(" Bảng sinh tự động")[0]
    return TableData(rows[0], body, used, source.rstrip("@").strip())


def source_table(rel_path: str, caption_key: str) -> TableData:
    """Bảng viết tay ngay trong nguồn báo cáo (`report/src`), ngay sau `@@CAP:tbl:<key>|…@@`."""
    lines = (ROOT / rel_path).read_text(encoding="utf-8").splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith(f"@@CAP:tbl:{caption_key}|"))
    block = []
    for line in lines[start + 1:]:
        if line.startswith("|"):
            block.append([c.strip() for c in line.strip().strip("|").split("|")])
        elif block:
            break
    note = next((line for line in lines[start:] if line.startswith("@@NOTE:")), "")
    source = note.removeprefix("@@NOTE:Nguồn:").rstrip("@").strip()
    return TableData(block[0], block[2:], [False] * (len(block) - 2), source)


def plain(value: str) -> str:
    return value.replace("**", "").replace("`", "").strip()


def number(value: str) -> float:
    return float(plain(value).replace(",", "."))


def column(data: TableData, name: str) -> list[float]:
    index = data.header.index(name)
    return [number(row[index]) for row in data.rows]


_NUMERIC = re.compile(r"^[\[\]\-−+\d.,%/ ()]+$|^—$")


def table(slide, data: TableData, x, y, w, *, widths=None, size=12, header=None, relabel=None,
          row_h=0.34, head_fill=NAVY):
    """Bảng: in đậm ô có **…**, nền xám hàng đang dùng, số căn phải."""
    head = header or data.header
    relabel = relabel or {}
    rows = [[relabel.get(plain(c), c.replace("`", "")) for c in row] for row in data.rows]
    shape = slide.shapes.add_table(len(rows) + 1, len(head), Inches(x), Inches(y), Inches(w),
                                   Inches(row_h * (len(rows) + 1)))
    tbl = shape.table
    tbl.horz_banding = False
    weights = widths or [1] * len(head)
    for index, weight in enumerate(weights):
        tbl.columns[index].width = Inches(w * weight / sum(weights))
    for r_index in range(len(rows) + 1):
        tbl.rows[r_index].height = Inches(row_h)
    for c_index, label in enumerate(head):
        _cell(tbl.cell(0, c_index), label, size, fill=head_fill, color=WHITE, bold=True,
              align=PP_ALIGN.CENTER)
    for r_index, row in enumerate(rows, start=1):
        fill = USED_FILL if data.used[r_index - 1] else WHITE
        for c_index, value in enumerate(row):
            bold = value.startswith("**") and value.endswith("**") and value.count("**") == 2
            value = value.replace("**", "")
            numeric = c_index > 0 and bool(_NUMERIC.match(value))
            _cell(tbl.cell(r_index, c_index), value, size, fill=fill, bold=bold,
                  align=PP_ALIGN.RIGHT if numeric else PP_ALIGN.LEFT)
    return shape


def _cell(cell, value, size, *, fill, color=TEXT, bold=False, align=PP_ALIGN.LEFT):
    cell.margin_left = cell.margin_right = Inches(0.06)
    cell.margin_top = cell.margin_bottom = Inches(0.025)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        line = etree.SubElement(tc_pr, qn(tag), w=str(Pt(0.75)))
        line_fill = etree.SubElement(line, qn("a:solidFill"))
        etree.SubElement(line_fill, qn("a:srgbClr")).set("val", "C9D2DD")
    solid = etree.SubElement(tc_pr, qn("a:solidFill"))
    etree.SubElement(solid, qn("a:srgbClr")).set("val", fill)


def source_note(slide, x, y, w, label) -> None:
    if isinstance(label, TableData):
        label = label.source
    text(slide, x, y, w, 0.3, "Nguồn: " + label, size=9.5, color=GRAY, italic=True)


def legend_note(slide, x, y, w) -> None:
    text(slide, x, y, w, 0.28, "In đậm: tốt nhất trong cột (không có nghĩa là khác biệt có ý "
         "nghĩa) · nền xám: cấu hình đang dùng", size=10, color=GRAY, italic=True)


# ---------------------------------------------------------------- biểu đồ
def bar_chart(slide, x, y, w, h, categories, series: dict[str, list[float]], *, colors,
              title_text, number_format="0.000", max_value=None):
    chart_data = CategoryChartData()
    chart_data.categories = categories
    for name, values in series.items():
        chart_data.add_series(name, values)
    graphic = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y),
                                     Inches(w), Inches(h), chart_data)
    chart = graphic.chart
    chart.has_title = True
    chart.chart_title.text_frame.text = title_text
    title_font = chart.chart_title.text_frame.paragraphs[0].runs[0].font
    title_font.size, title_font.bold, title_font.name = Pt(13), True, FONT
    title_font.color.rgb = rgb(NAVY)
    chart.has_legend = len(series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size, chart.legend.font.name = Pt(11), FONT
    plot = chart.plots[0]
    plot.gap_width = 70
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format, labels.number_format_is_linked = number_format, False
    labels.font.size, labels.font.name = Pt(11), FONT
    labels.font.color.rgb = rgb(TEXT)
    for index, series_obj in enumerate(plot.series):
        series_obj.format.fill.solid()
        series_obj.format.fill.fore_color.rgb = rgb(colors[index])
    value_axis = chart.value_axis
    value_axis.has_major_gridlines = True
    value_axis.major_gridlines.format.line.color.rgb = rgb(LINE)
    value_axis.tick_labels.font.size, value_axis.tick_labels.font.name = Pt(10), FONT
    value_axis.tick_labels.font.color.rgb = rgb(GRAY)
    value_axis.format.line.fill.background()
    value_axis.minimum_scale = 0
    if max_value is not None:
        value_axis.maximum_scale = max_value
    category_axis = chart.category_axis
    category_axis.tick_labels.font.size, category_axis.tick_labels.font.name = Pt(11), FONT
    category_axis.tick_labels.font.color.rgb = rgb(TEXT)
    category_axis.format.line.color.rgb = rgb(LINE)
    return graphic
