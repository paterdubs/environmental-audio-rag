"""Kiểm thử bộ dựng slide giữa kỳ: đọc bảng, đánh dấu đậm, và bộ slide hoàn chỉnh."""

from __future__ import annotations

from pathlib import Path

import pytest

pytest.importorskip("pptx")

from pptx import Presentation  # noqa: E402

from scripts.midterm_slides import kit  # noqa: E402

SOURCE = """# Chương

@@CAP:tbl:plan|Kế hoạch@@

| Thời gian | Việc |
|---|---|
| 04/10 | Viết **báo cáo** |
| 05/10 | Làm slide |

@@NOTE:Nguồn: docs/PLAN.md.@@
"""


def test_source_table_reads_rows_after_caption(tmp_path: Path, monkeypatch):
    path = tmp_path / "ch.md"
    path.write_text(SOURCE, encoding="utf-8")
    monkeypatch.setattr(kit, "ROOT", tmp_path)

    data = kit.source_table("ch.md", "plan")

    assert data.header == ["Thời gian", "Việc"]
    assert data.rows == [["04/10", "Viết **báo cáo**"], ["05/10", "Làm slide"]]
    assert data.source == "docs/PLAN.md."


def test_table_bolds_whole_cells_and_strips_inline_markers():
    pres = Presentation()
    slide = pres.slides.add_slide(pres.slide_layouts[6])
    data = kit.TableData(["A", "B"], [["**0.5**", "Viết **chỉ trên dev**"]], [True], "")

    shape = kit.table(slide, data, 0.5, 0.5, 6.0)

    first, second = shape.table.cell(1, 0), shape.table.cell(1, 1)
    assert first.text == "0.5" and first.text_frame.paragraphs[0].runs[0].font.bold
    assert second.text == "Viết chỉ trên dev"
    assert not second.text_frame.paragraphs[0].runs[0].font.bold


ARTIFACT = kit.ROOT / "ml/runs/sed_ensemble_s13_f2_20260929T045053Z/evaluation_cv_annotated.json"


@pytest.mark.skipif(not ARTIFACT.exists(), reason="cần artifact đã khóa trong ml/runs")
def test_deck_has_numbered_slides_and_clean_text(tmp_path: Path):
    from scripts.build_midterm_slides import build

    deck = build(tmp_path / "deck.pptx")

    assert (deck.count, deck.appendix) == (29, 7)
    pres = Presentation(str(tmp_path / "deck.pptx"))
    for slide in pres.slides:
        texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
        texts += [cell.text for shape in slide.shapes if shape.has_table
                  for row in shape.table.rows for cell in row.cells]
        joined = " ".join(texts)
        assert "**" not in joined and "`" not in joined
        assert "Tương tác Người" not in joined
